"""Export assembled markdown manuscript to DOCX format."""

import re
from pathlib import Path

try:
    from docx import Document
    from docx.shared import Pt, Inches
    from docx.enum.text import WD_ALIGN_PARAGRAPH

    HAS_DOCX = True
except ImportError:
    HAS_DOCX = False


class DocxExporter:
    """Convert a markdown manuscript string to a formatted Word document.

    Follows mass-market fiction conventions: Times New Roman 12pt,
    1.5 line spacing, first paragraph no indent, subsequent 0.5" indent.
    """

    def __init__(self, concept_seed: dict = None):
        self.concept_seed = concept_seed or {}

    def export(
        self,
        manuscript_md: str,
        output_path: str = "output/_fallback/export/manuscript.docx",
    ) -> Path:
        """Convert markdown manuscript to DOCX.

        Args:
            manuscript_md: The assembled markdown manuscript string.
            output_path: Where to save the DOCX file.

        Returns:
            Path to the created file.
        """
        if not HAS_DOCX:
            raise ImportError(
                "python-docx is required for DOCX export. "
                "Install with: pip install python-docx>=1.1.0"
            )

        doc = Document()
        self._apply_defaults(doc)
        self._add_title_page(doc)

        chapters = self._parse_chapters(manuscript_md)
        for i, chapter in enumerate(chapters):
            self._add_chapter(doc, chapter["number"], chapter["content"], is_first=(i == 0))

        output = Path(output_path)
        output.parent.mkdir(parents=True, exist_ok=True)
        doc.save(str(output))
        return output

    def _apply_defaults(self, doc: "Document") -> None:
        """Set document-wide defaults."""
        style = doc.styles["Normal"]
        font = style.font
        font.name = "Times New Roman"
        font.size = Pt(12)
        paragraph_format = style.paragraph_format
        paragraph_format.line_spacing = 1.5

        # Set margins
        for section in doc.sections:
            section.top_margin = Inches(1)
            section.bottom_margin = Inches(1)
            section.left_margin = Inches(1)
            section.right_margin = Inches(1)

    def _add_title_page(self, doc: "Document") -> None:
        """Add a centered title page."""
        meta = self.concept_seed.get("meta", {})
        title = meta.get("project_title", "Untitled Manuscript")
        franchise = meta.get("franchise", "")

        # Add spacing before title
        for _ in range(6):
            doc.add_paragraph("")

        title_para = doc.add_paragraph()
        title_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = title_para.add_run(title)
        run.bold = True
        run.font.size = Pt(24)
        run.font.name = "Times New Roman"

        if franchise:
            subtitle = doc.add_paragraph()
            subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = subtitle.add_run(f"A {franchise} Novel")
            run.italic = True
            run.font.size = Pt(14)
            run.font.name = "Times New Roman"

        # Page break after title
        doc.add_page_break()

    def _parse_chapters(self, manuscript_md: str) -> list[dict]:
        """Parse markdown into chapter blocks."""
        chapters = []
        # Split on # Chapter N headings
        parts = re.split(r"^# Chapter (\d+)\s*$", manuscript_md, flags=re.MULTILINE)

        # parts[0] is front matter, then alternating (number, content)
        for i in range(1, len(parts) - 1, 2):
            chapter_num = int(parts[i])
            content = parts[i + 1].strip()
            chapters.append({"number": chapter_num, "content": content})

        return chapters

    def _add_chapter(
        self, doc: "Document", chapter_num: int, content: str, is_first: bool
    ) -> None:
        """Add a single chapter to the document."""
        if not is_first:
            doc.add_page_break()

        # Chapter heading
        heading = doc.add_heading(f"Chapter {chapter_num}", level=1)
        heading.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for run in heading.runs:
            run.font.name = "Times New Roman"
            run.font.size = Pt(18)

        # Split content by scene breaks
        scenes = re.split(r"\n\s*\*\s*\*\s*\*\s*\n", content)

        for scene_idx, scene_text in enumerate(scenes):
            if scene_idx > 0:
                # Scene break marker
                break_para = doc.add_paragraph()
                break_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
                break_para.paragraph_format.space_before = Pt(18)
                break_para.paragraph_format.space_after = Pt(18)
                run = break_para.add_run("* * *")
                run.font.name = "Times New Roman"
                run.font.size = Pt(12)

            paragraphs = [p.strip() for p in scene_text.strip().split("\n\n") if p.strip()]
            for p_idx, para_text in enumerate(paragraphs):
                # Skip any remaining markdown headings or horizontal rules
                if para_text.startswith("#") or para_text == "---":
                    continue

                p = doc.add_paragraph(para_text)
                p.style = doc.styles["Normal"]

                # First paragraph of chapter/scene: no indent
                # Subsequent paragraphs: 0.5" first-line indent
                if p_idx > 0:
                    p.paragraph_format.first_line_indent = Inches(0.5)
                else:
                    p.paragraph_format.first_line_indent = Inches(0)
